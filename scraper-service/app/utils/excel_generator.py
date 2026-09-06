import io
from typing import List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import JobPost

def generate_excel_bytes(jobs: List[JobPost]) -> io.BytesIO:
    """
    Generates a beautifully styled Excel workbook containing all jobs
    with active hyperlinks, match score highlights, and fitted columns.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Job Openings"

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # 1. Define Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark Navy Slate
    
    cell_font = Font(name="Calibri", size=10, color="1E293B")
    link_font = Font(name="Calibri", size=10, color="2563EB", underline="single") # Blue hyperlink
    
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )
    
    # Score fills
    high_match_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Light Emerald Green
    high_match_font = Font(name="Calibri", size=10, bold=True, color="166534") # Dark Green
    
    mid_match_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid") # Light Amber Yellow
    mid_match_font = Font(name="Calibri", size=10, bold=True, color="854D0E") # Dark Amber

    # 2. Header Columns
    headers = [
        "Match Score",
        "Job Title",
        "Company Name",
        "Company Details",
        "Location",
        "Required Skills",
        "Experience",
        "Salary",
        "Date Posted",
        "Job / Apply URL",
        "Apply Method",
        "Recruiter Name",
        "Recruiter Email",
        "Recruiter Phone",
        "Source Website",
        "Match Summary",
        "Status"
    ]

    ws.append(headers)

    # Style Header Row
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # 3. Insert Data Rows
    for row_idx, job in enumerate(jobs, start=2):
        skills_str = ", ".join(job.required_skills) if job.required_skills else ""
        score_val = f"{job.match_score}%" if job.match_score is not None else "N/A"
        
        row_data = [
            score_val,
            job.title or "",
            job.company or "",
            job.company_details or "",
            job.location or "",
            skills_str,
            job.experience or "",
            job.salary or "Not Disclosed",
            job.date_posted or "",
            job.job_url or "",
            job.apply_method or "Direct Link",
            job.recruiter_name or "",
            job.recruiter_email or "",
            job.recruiter_phone or "",
            job.source_website or "",
            job.match_summary or "",
            job.status or "New"
        ]
        ws.append(row_data)

        # Style Data Cells
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = cell_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

            # Center-align specific columns
            if col_idx in [1, 5, 7, 8, 9, 11, 14, 15, 17]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Highlight Match Score
            if col_idx == 1 and job.match_score is not None:
                if job.match_score >= 80:
                    cell.fill = high_match_fill
                    cell.font = high_match_font
                elif job.match_score >= 60:
                    cell.fill = mid_match_fill
                    cell.font = mid_match_font

            # Hyperlink for Job URL
            if col_idx == 10 and job.job_url and job.job_url.startswith("http"):
                cell.hyperlink = job.job_url
                cell.font = link_font
                cell.value = "Open Job Page ↗"

            # Hyperlink for Recruiter Email
            if col_idx == 13 and job.recruiter_email and "@" in job.recruiter_email:
                cell.hyperlink = f"mailto:{job.recruiter_email}?subject=Application for {job.title} position"
                cell.font = link_font

        ws.row_dimensions[row_idx].height = 22

    # 4. Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        # Add buffer and clamp between 12 and 45
        ws.column_dimensions[col_letter].width = max(min(max_len + 3, 45), 12)

    # Specific overrides for readability
    ws.column_dimensions["A"].width = 14  # Match Score
    ws.column_dimensions["B"].width = 30  # Title
    ws.column_dimensions["C"].width = 25  # Company
    ws.column_dimensions["E"].width = 20  # Location
    ws.column_dimensions["F"].width = 35  # Skills
    ws.column_dimensions["J"].width = 18  # Job URL link
    ws.column_dimensions["P"].width = 40  # Match Summary

    # Freeze top row
    ws.freeze_panes = "A2"

    # Enable filter dropdowns
    ws.auto_filter.ref = ws.dimensions

    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
